#!/usr/bin/env python3
"""
Exporte la feuille de calcul vers un jeu prêt pour le math SDK de Stake Engine.

    python exporter.py ../modele_maths/modele_maths.xlsx
    python exporter.py ma_feuille.xlsx --game-id 12_3_monjeu --nom "Mon Jeu"

Écrit dans games/<game-id>/ :
  - params.json      : paytable, lignes, multis, free spins, bonus buy, cibles d'optimisation
  - reels/BR0.csv    : bandes du jeu de base (poids de l'onglet Poids = nombre d'occurrences)
  - reels/FR0.csv    : bandes des free spins
  - reels/FRWCAP.csv : bandes FS enrichies en globes, utilisées seulement pour fabriquer
                       des résultats « max win » (critère wincap de l'optimiseur)
Les fichiers Python du jeu (game_config.py, gamestate.py…) sont recopiés depuis le
modèle games/0_0_globe s'ils n'existent pas encore.
"""
import argparse
import json
import random
import shutil
import sys
from pathlib import Path

ICI = Path(__file__).resolve().parent
sys.path.insert(0, str(ICI.parent / "modele_maths"))
import simulateur  # noqa: E402  (lecture des plages nommées du classeur)

MODELE = ICI / "games" / "0_0_globe"
FICHIERS_JEU = ["game_config.py", "game_calculations.py", "game_executables.py", "game_override.py",
                "game_events.py", "gamestate.py", "game_optimization.py", "run.py", "readme.txt"]
COMBOS_SDK = {("CLUSTER", "WILD_ADD"), ("CLUSTER", "GLOBAL_SUM"), ("LINES", "WILD_ADD"), ("LINES", "GLOBAL_SUM"), ("WAYS", "WILD_MULT"), ("WAYS", "GLOBAL_SUM")}
RESERVES = {"W", "MX"}
WINCAP_BOOST_GLOBE = 10   # globes x10 sur la bande FRWCAP


def erreur(msg):
    sys.exit(f"Erreur : {msg}")


def lire_simulation(path):
    """Résultats de simulateur.py (onglet Simulation, cellules constantes)."""
    from openpyxl import load_workbook
    ws = load_workbook(path)["Simulation"]
    v = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value for r in range(5, 23)}
    need = ["RTP jeu de base (globe inclus)", "RTP free spins", "RTP total", "Hit frequency", "Bonus 1 spin sur"]
    if any(not isinstance(v.get(k), (int, float)) for k in need):
        erreur("l'onglet Simulation est vide : lance d'abord  python simulateur.py <classeur>")
    return v


def comptes(poids, L):
    """Répartit L cases proportionnellement aux poids (plus forts restes), au moins 1 si poids > 0."""
    tot = float(sum(poids))
    brut = [w * L / tot for w in poids]
    n = [max(int(b), 1 if w > 0 else 0) for b, w in zip(brut, poids)]
    ordre = sorted(range(len(poids)), key=lambda i: brut[i] - int(brut[i]), reverse=True)
    i = 0
    while sum(n) < L:
        if poids[ordre[i % len(ordre)]] > 0:
            n[ordre[i % len(ordre)]] += 1
        i += 1
    while sum(n) > L:                      # cas rare : trop de minimums à 1
        j = max(range(len(n)), key=lambda k: n[k])
        n[j] -= 1
    return n


def bande(poids, codes, L, ecart_scatter, scatter, rng):
    """Construit une bande de L cases, chaque symbole en proportion de son poids.
    Les symboles bonus sont espacés d'au moins `ecart_scatter` cases (au plus 1 symbole bonus visible
    par rouleau), comme l'exige force_special_board du SDK."""
    n = dict(zip(codes, comptes(list(poids), L)))
    nb_sc = n.pop(scatter, 0)
    reste = [c for c, k in n.items() for _ in range(k)]
    rng.shuffle(reste)
    L = len(reste) + nb_sc
    if nb_sc and L < nb_sc * ecart_scatter:
        erreur(f"trop de symboles bonus ({nb_sc}) pour une bande de {L} cases : il faut au moins "
               f"{ecart_scatter} cases par symbole bonus.")
    # Répartit les symboles bonus avec un écart >= ecart_scatter (bande circulaire)
    libre = L - nb_sc * ecart_scatter
    coupes = sorted(rng.randint(0, libre) for _ in range(nb_sc))
    pos = [c + i * ecart_scatter for i, c in enumerate(coupes)]
    strip, it = [], iter(reste)
    pos_set = set(pos)
    for i in range(L):
        strip.append(scatter if i in pos_set else next(it))
    return strip


def ecrire_bandes(path, bandes):
    """Format CSV du SDK : une colonne par rouleau, toutes de même longueur."""
    assert len({len(b) for b in bandes}) == 1
    with open(path, "w", encoding="utf-8") as f:
        for i in range(len(bandes[0])):
            f.write(",".join(b[i] for b in bandes) + "\n")


def proba_bonus(bandes, scatter, rows):
    """Loi exacte du nombre de symboles bonus visibles (au plus 1 par rouleau)."""
    dist = [1.0]
    for b in bandes:
        q = rows * b.count(scatter) / len(b)
        new = [0.0] * (len(dist) + 1)
        for k, pk in enumerate(dist):
            new[k] += pk * (1 - q)
            new[k + 1] += pk * q
        dist = new
    return dist


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("classeur")
    ap.add_argument("--game-id", default="0_0_globe", help="dossier/ID du jeu (ex. 12_3_monjeu)")
    ap.add_argument("--nom", default="Globe Multipliers", help="nom de travail du jeu")
    ap.add_argument("--provider", type=int, default=0, help="numéro de provider Stake Engine")
    ap.add_argument("--seed", type=int, default=2024, help="graine pour mélanger les bandes")
    a = ap.parse_args()

    p = simulateur.load_params(a.classeur)
    sim = lire_simulation(a.classeur)
    rows = p["rows"]
    gain, mode = p["mode_gain"], p["mode"]
    if (gain, mode) not in COMBOS_SDK:
        erreur(f"{gain} + {mode} n'existe pas nativement dans le math SDK. Combinaisons possibles : "
               + ", ".join(f"{g} + {m}" for g, m in sorted(COMBOS_SDK)))
    codes, types = p["codes"], p["types"]
    for c in codes:
        if not str(c).isalnum() or c in RESERVES:
            erreur(f"code symbole invalide « {c} » : lettres/chiffres uniquement, et pas W ni MX (réservés).")
    scatter = codes[types.index("SCATTER")]
    if "GLOBE" not in types:
        erreur("il faut un symbole de type GLOBE.")
    globe = codes[types.index("GLOBE")]
    pay_codes = [c for c, t in zip(codes, types) if t == "PAY"]
    table = dict(p["table_fs"])
    if min(table) != p["scatter_min"]:
        erreur("SCATTER_MIN doit être égal au premier seuil de la table des FS.")
    if max(table) > 5:
        erreur("au plus 1 symbole bonus par rouleau dans le SDK : la table des FS ne peut pas dépasser 5 symboles bonus.")
    if p["fs_par_bonus"] != int(p["fs_par_bonus"]) or p["fs_par_bonus"] < 1:
        erreur("FS_PAR_BONUS doit être un entier >= 1.")
    for nom in ("poids_base", "poids_fs"):
        w = p[nom]
        if (abs(w - w.round()) > 1e-9).any():
            erreur(f"{nom} : les poids doivent être des entiers (= nombre d'occurrences sur la bande).")

    dossier = ICI / "games" / a.game_id
    (dossier / "reels").mkdir(parents=True, exist_ok=True)
    if dossier != MODELE:
        for f in FICHIERS_JEU:
            if not (dossier / f).exists():
                shutil.copy(MODELE / f, dossier / f)

    rng = random.Random(a.seed)

    def bandes(poids):
        L = int(poids.sum(axis=0).max())   # même longueur pour les 5 rouleaux
        return [bande(poids[:, r], codes, L, rows, scatter, rng) for r in range(5)]

    br = bandes(p["poids_base"])
    fr = bandes(p["poids_fs"])
    wcap_poids = p["poids_fs"].copy()
    wcap_poids[codes.index(globe)] *= WINCAP_BOOST_GLOBE
    wc = bandes(wcap_poids)
    ecrire_bandes(dossier / "reels" / "BR0.csv", br)
    ecrire_bandes(dossier / "reels" / "FR0.csv", fr)
    ecrire_bandes(dossier / "reels" / "FRWCAP.csv", wc)

    # Probabilités naturelles de 3/4/5 symboles bonus : servent à forcer le bonus avec les bonnes proportions
    dist = proba_bonus(br, scatter, rows)
    seuils = sorted(table)
    trig = {}
    for i, s in enumerate(seuils):
        hi = seuils[i + 1] if i + 1 < len(seuils) else 6
        trig[str(s)] = round(sum(dist[s:hi]), 10)

    # Cibles d'optimisation tirées de la simulation de la feuille
    rtp = round(p["rtp_cible"], 4)
    rtp_wincap = 0.001
    part_fs = sim["RTP free spins"] / sim["RTP total"]
    rtp_fs = round((rtp - rtp_wincap) * part_fs, 4)
    rtp_bg = round(rtp - rtp_wincap - rtp_fs, 4)
    mv = lambda poids: {str(int(v) if float(v).is_integer() else v): float(w)
                        for v, w in zip(p["multi_val"], poids) if w > 0}
    wcap_mult = p["multi_w_fs"] * p["multi_val"]          # multis plus forts pour les résultats max win

    if gain == "CLUSTER":
        # Une entrée par taille de cluster (kind = taille), comme convert_range_table du SDK
        n_cases = rows * 5
        paytable = []
        for c in pay_codes:
            i = codes.index(c)
            for j, t in enumerate(p["tailles"]):
                fin = p["tailles"][j + 1] - 1 if j + 1 < len(p["tailles"]) else n_cases
                for k in range(t, min(fin, n_cases) + 1):
                    if p["pays_cluster"][i][j] > 0:
                        paytable.append({"symbol": c, "kind": k, "pay": float(p["pays_cluster"][i][j])})
    else:
        paytable = [{"symbol": c, "kind": k, "pay": float(p["pays"][codes.index(c)][k - 3])}
                    for c in pay_codes for k in (3, 4, 5) if p["pays"][codes.index(c)][k - 3] > 0]

    params = {
        "_info": "Généré par exporter.py depuis la feuille de calcul. Ne pas éditer à la main : "
                 "modifie la feuille puis relance l'export.",
        "game_id": a.game_id,
        "provider_number": a.provider,
        "working_name": a.nom,
        "rtp": rtp,
        "wincap": p["max_win"],
        "win_type": gain.lower(),
        "multi_mode": mode,
        "num_rows": rows,
        "num_reels": 5,
        "paytable": paytable,
        "cascades": bool(p["cascades"]) if gain == "CLUSTER" else False,
        "top_symbol": pay_codes[0],
        "scatter": scatter,
        "globe": globe,
        "globe_multiplier": p["multi_globe"],
        "paylines": p["lignes"].tolist(),
        "fs_triggers_base": {str(k): v for k, v in table.items()},
        "fs_per_scatter_in_fs": int(p["fs_par_bonus"]),
        "scatter_trigger_weights": trig,
        "mult_values": {"basegame": mv(p["multi_w_base"]), "freegame": mv(p["multi_w_fs"]),
                        "wincap": mv(wcap_mult)},
        "buy": {"cost": p["buy_cout"], "spins": p["buy_fs"]},
        "targets": {
            "wincap_rtp": rtp_wincap,
            "freegame_rtp": rtp_fs,
            "freegame_hr": round(float(sim["Bonus 1 spin sur"]), 2),
            "basegame_rtp": rtp_bg,
            "basegame_hr": round(1 / float(sim["Hit frequency"]), 3),
        },
    }
    with open(dossier / "params.json", "w", encoding="utf-8") as f:
        json.dump(params, f, ensure_ascii=False, indent=2)

    print(f"Jeu exporté dans {dossier}")
    detail = {"LINES": f"{len(params['paylines'])} lignes", "WAYS": f"{rows ** 5} ways",
              "CLUSTER": f"clusters de {p['tailles'][0]}+, cascades {'oui' if params['cascades'] else 'non'}"}[gain]
    print(f"  mode {gain} + {mode}, {detail}, max win {p['max_win']:g}x, RTP cible {rtp:.2%}")
    print(f"  bandes : base {len(br[0])} cases, FS {len(fr[0])} cases ; "
          f"bonus naturel 1 spin sur {1 / sum(dist[min(table):]):.0f}")
    print(f"  cibles : base {rtp_bg:.4f} (hr {params['targets']['basegame_hr']}), "
          f"FS {rtp_fs:.4f} (1/{params['targets']['freegame_hr']}), wincap {rtp_wincap}")


if __name__ == "__main__":
    main()
